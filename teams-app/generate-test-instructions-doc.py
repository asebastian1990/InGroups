#!/usr/bin/env python3
"""Generate TEAMS-STORE-TEST-INSTRUCTIONS.docx for Partner Center submission."""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

OUT = Path(__file__).resolve().parent / "TEAMS-STORE-TEST-INSTRUCTIONS.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_assumption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run("ASSUMPTION: ")
    run.bold = True
    run.font.color.rgb = RGBColor(0x8B, 0x45, 0x13)
    p.add_run(text)


def add_media_placeholder(doc: Document, kind: str, description: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"[INSERT {kind.upper()}: {description}]")
    run.italic = True
    run.font.color.rgb = RGBColor(0x00, 0x55, 0xAA)


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("InGroups — Teams Store validation test instructions", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    add_body(
        doc,
        "Document version: 1.0 (app package version 1.0.7). "
        "Prepared for Microsoft Teams Store certification (Partner Center test notes).",
    )
    doc.add_paragraph()

    # Table of contents (manual, like Sample 1)
    add_heading(doc, "Table of Contents", 1)
    for item in [
        "About the app",
        "Pre-requisites",
        "Test credentials",
        "App functionality",
        "  Landing (Teams personal tab and web)",
        "  Host a game lobby",
        "  Join from Microsoft Teams",
        "  Join from web browser (guest or signed-in)",
        "  Start a game and play one round",
        "  Premium word sets (licensed host)",
        "  License purchase and activation (optional path)",
        "  Meeting side panel and channel tab",
        "  Privacy policy and terms of use",
        "Demo video",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_page_break()

    add_heading(doc, "About the app", 1)
    add_body(
        doc,
        "InGroups is a facilitated word-game icebreaker for meetings, workshops, and team gatherings. "
        "Players align on a secret word from a shared grid without giving it away. Each round, one group "
        "becomes the In Group and discusses openly; Out Groups observe and try to predict the chosen word. "
        "Optional timers, scoring, and win conditions keep sessions moving.",
    )
    add_body(
        doc,
        "InGroups runs as a Microsoft Teams personal tab, meeting side panel, and channel/chat tab, "
        "and as a standalone web app at the same production URL. Multiplayer sessions use a four-letter "
        "room code; players in Teams and players in a browser can join the same lobby.",
    )
    add_body(
        doc,
        "A one-time license (Stripe checkout) unlocks 50+ premium word sets and custom word-set creation "
        "for the purchasing account. Free tier includes built-in word sets for hosting and play.",
    )

    add_heading(doc, "Pre-requisites", 1)
    add_body(
        doc,
        "IMPORTANT: Before testing, ensure the following:",
    )
    add_bullet(doc, "Microsoft Teams desktop or web client (current channel).")
    add_bullet(
        doc,
        "Network access to https://ingroups.annaliese-sebastian.com and "
        "https://ingroups-production.up.railway.app (API).",
    )
    add_bullet(
        doc,
        "At least three participants in the same room to start a game "
        "(MIN_PLAYERS = 3 in the product).",
    )
    add_bullet(
        doc,
        "Install or open InGroups from the Teams app catalog entry under validation, "
        "or sideload package version 1.0.7 if testing before Store listing is live.",
    )
    add_assumption(
        doc,
        "No separate Global Administrator / Teams Administrator test account is supplied. "
        "Validation should use user-level access: install/open InGroups from Teams Apps for the "
        "provided non-admin accounts, or open the web app directly. In tenants that block "
        "third-party apps by policy, a tenant admin may need to allow InGroups; that admin "
        "workflow is not included in the credentials below.",
    )
    add_assumption(
        doc,
        "Test accounts live in the publisher’s Microsoft 365 tenant (domain and passwords "
        "filled in below before submission). Replace placeholder emails/passwords with real "
        "validation-only users.",
    )
    add_assumption(
        doc,
        "Account tester1@ has an active InGroups license pre-provisioned on the production "
        "backend before review begins. tester2@ and tester3@ do not require a license to "
        "join and play in a lobby hosted by tester1@.",
    )
    add_body(
        doc,
        "External services used during testing: Microsoft Entra ID (Teams SSO), Clerk "
        "(account/session for signed-in users), Stripe (license purchase if testing checkout), "
        "Railway-hosted API and WebSockets (real-time game state).",
    )

    add_heading(doc, "Test credentials", 1)
    add_body(doc, "Use these accounts only for Microsoft validation.")
    doc.add_paragraph()

    cred_table = doc.add_table(rows=4, cols=3)
    cred_table.style = "Table Grid"
    headers = ("Role", "Email", "Password")
    for i, h in enumerate(headers):
        cred_table.rows[0].cells[i].text = h
    rows = [
        (
            "Host (licensed account — premium word sets)",
            "tester1@YOUR-TENANT-DOMAIN",
            "REPLACE-BEFORE-SUBMIT",
        ),
        (
            "Player 2 (Teams or web)",
            "tester2@YOUR-TENANT-DOMAIN",
            "REPLACE-BEFORE-SUBMIT",
        ),
        (
            "Player 3 / first-run (has not opened InGroups before review)",
            "tester3@YOUR-TENANT-DOMAIN",
            "REPLACE-BEFORE-SUBMIT",
        ),
    ]
    for r, (role, email, pw) in enumerate(rows, start=1):
        cred_table.rows[r].cells[0].text = role
        cred_table.rows[r].cells[1].text = email
        cred_table.rows[r].cells[2].text = pw

    doc.add_paragraph()
    add_body(
        doc,
        "Additional players (no Microsoft account required): open "
        "https://ingroups.annaliese-sebastian.com in a browser, choose Continue as Guest, "
        "enter a display name and the host’s four-letter room code. Multiple guest browsers "
        "can join the same lobby as the Teams users.",
    )
    add_assumption(
        doc,
        "YOUR-TENANT-DOMAIN is replaced with your actual test tenant domain "
        "(e.g. contoso.onmicrosoft.com) before Partner Center submission.",
    )

    add_heading(doc, "App functionality", 1)

    add_heading(doc, "Landing (Teams personal tab and web)", 2)
    add_body(doc, "Teams personal tab URL: https://ingroups.annaliese-sebastian.com/teams")
    add_body(doc, "Web app URL: https://ingroups.annaliese-sebastian.com/")
    add_body(
        doc,
        "Steps — Teams (tester1@):",
    )
    add_bullet(doc, "Open Microsoft Teams → Apps → InGroups (or Preview in Teams from Developer Portal).")
    add_bullet(
        doc,
        "On first load, the app connects to Teams and signs in via Microsoft Teams SSO "
        "(Entra) into a Clerk-backed session when SSO succeeds.",
    )
    add_bullet(
        doc,
        "If prompted, complete sign-in; alternatively Continue as Guest is available "
        "for gameplay without a Clerk account (license features require sign-in).",
    )
    add_bullet(
        doc,
        "Landing shows Host New Game, room code join, View Word Sets, Create Word Set, License, "
        "and links to Privacy Policy and Terms of Use.",
    )
    add_body(doc, "Steps — Web (tester3@ first-run OR guest browser):")
    add_bullet(doc, "Navigate to https://ingroups.annaliese-sebastian.com/")
    add_bullet(
        doc,
        "First-time signed-in users: use Sign in (Google or email) or Continue as Guest.",
    )
    add_bullet(doc, "Same landing actions as Teams: host, join, word sets, license.")
    add_media_placeholder(
        doc,
        "screenshot",
        "Teams personal tab — InGroups landing with Host New Game and room code field.",
    )
    add_media_placeholder(
        doc,
        "screenshot",
        "Web landing at ingroups.annaliese-sebastian.com (signed-in and/or Continue as Guest).",
    )

    add_heading(doc, "Host a game lobby", 2)
    add_body(doc, "Steps (tester1@ — licensed host):")
    add_bullet(doc, "Enter display name if empty → click Host New Game.")
    add_bullet(doc, "Lobby shows a four-letter Room code banner (share with other players).")
    add_bullet(
        doc,
        "Host can adjust Points to Win, open View Word Sets, and see the player list.",
    )
    add_bullet(
        doc,
        "Start Game stays disabled until at least three players have joined the lobby.",
    )
    add_media_placeholder(
        doc,
        "screenshot",
        "Host lobby with room code visible and player list (1–2 players waiting).",
    )

    add_heading(doc, "Join from Microsoft Teams", 2)
    add_body(doc, "Steps (tester2@):")
    add_bullet(doc, "Open InGroups in Teams (personal tab).")
    add_bullet(doc, "Enter display name → enter the host’s room code → Join Game.")
    add_bullet(doc, "Confirm name appears in the host’s player list.")
    add_body(doc, "Steps (tester3@ — first-run):")
    add_bullet(
        doc,
        "Do not open InGroups before this test if you need a clean first-run experience.",
    )
    add_bullet(doc, "Open InGroups in Teams → complete SSO or guest path → join with room code.")
    add_media_placeholder(
        doc,
        "screenshot",
        "Join Game with room code filled in (Teams client).",
    )

    add_heading(doc, "Join from web browser (guest or signed-in)", 2)
    add_body(
        doc,
        "This demonstrates cross-client play: Teams host + web joiners in one lobby.",
    )
    add_bullet(
        doc,
        "In Chrome/Edge, open https://ingroups.annaliese-sebastian.com (incognito optional).",
    )
    add_bullet(doc, "Click Continue as Guest (or sign in with a separate test account).")
    add_bullet(doc, "Enter a display name and the same room code → Join Game.")
    add_bullet(
        doc,
        "Repeat in a second browser profile if you need more than three total players quickly.",
    )
    add_assumption(
        doc,
        "Validation team may use guest web join instead of tester3@ if three distinct "
        "signed-in Teams accounts are not available; the product supports both.",
    )
    add_media_placeholder(
        doc,
        "screenshot",
        "Web guest joined to the same lobby as Teams users (player list shows 3+ names).",
    )

    add_heading(doc, "Start a game and play one round", 2)
    add_body(doc, "Steps (after ≥3 players in lobby):")
    add_bullet(doc, "Host (tester1@): optionally open View Word Sets and select a free word set.")
    add_bullet(doc, "Host: click Start Game.")
    add_bullet(
        doc,
        "Assign groups / In Group vs Out Groups per on-screen host controls; start a round.",
    )
    add_bullet(
        doc,
        "In Group selects a word from the grid; Out Groups submit guesses before time expires.",
    )
    add_bullet(doc, "Complete at least one scoring round; host may start another or end session.")
    add_media_placeholder(doc, "screenshot", "Active gameplay — word grid and group assignment.")
    add_media_placeholder(doc, "screenshot", "Round end / scores visible.")

    add_heading(doc, "Premium word sets (licensed host)", 2)
    add_body(doc, "Steps (tester1@ only — pre-licensed):")
    add_bullet(doc, "From landing or lobby → View Word Sets.")
    add_bullet(
        doc,
        "Confirm Premium-tagged sets are selectable when hosting (not locked to purchase).",
    )
    add_bullet(doc, "Host a lobby, select a premium word set, start game with tester2@ / tester3@.")
    add_assumption(
        doc,
        "If premium sets appear locked on tester1@, contact support@ (see below) — "
        "license should be active on production before review.",
    )
    add_media_placeholder(
        doc,
        "screenshot",
        "View Word Sets showing premium sets available to licensed host.",
    )

    add_heading(doc, "License purchase and activation (optional path)", 2)
    add_body(
        doc,
        "Primary validation path uses tester1@ with a pre-granted license. To test checkout:",
    )
    add_bullet(
        doc,
        "Sign in on web or Teams (not guest) with an account without a license.",
    )
    add_bullet(doc, "Landing → License → purchase via Stripe Checkout (live mode in production).")
    add_bullet(
        doc,
        "After payment, return URL https://ingroups.annaliese-sebastian.com/purchase/complete "
        "confirms fulfillment; License screen shows active license or unused keys.",
    )
    add_bullet(doc, "Activate a license key on License screen if testing key-based activation.")
    add_assumption(
        doc,
        "Microsoft validators are not required to complete a live purchase if tester1@ "
        "already demonstrates premium content; include purchase steps only if testing monetization.",
    )
    add_media_placeholder(doc, "screenshot", "License screen with active license summary.")

    add_heading(doc, "Meeting side panel and channel tab", 2)
    add_body(
        doc,
        "InGroups supports meeting side panel, meeting chat tab, channel tab, and private chat tab "
        "(see manifest static/configurable tabs).",
    )
    add_bullet(
        doc,
        "In a Teams meeting → Apps → InGroups → pin/open side panel (or add via + Apps).",
    )
    add_bullet(
        doc,
        "Same /teams experience: host or join with room code; meeting participants can share one lobby.",
    )
    add_bullet(
        doc,
        "Optional: Team/channel → + Add a tab → InGroups → save configuration "
        "(https://ingroups.annaliese-sebastian.com/teams/config).",
    )
    add_media_placeholder(
        doc,
        "screenshot",
        "InGroups open in a Teams meeting side panel with lobby or game visible.",
    )

    add_heading(doc, "Privacy policy and terms of use", 2)
    add_bullet(doc, "Privacy: https://ingroups.annaliese-sebastian.com/privacypolicy")
    add_bullet(doc, "Terms: https://ingroups.annaliese-sebastian.com/termsofuse")
    add_bullet(doc, "Also linked from the InGroups landing footer in Teams and web.")
    add_media_placeholder(
        doc,
        "screenshot",
        "Privacy policy page loads without sign-in.",
    )

    add_heading(doc, "Demo video", 1)
    add_body(
        doc,
        "Provide a public/unlisted video URL in Partner Center (YouTube or Vimeo per Microsoft guidelines).",
    )
    add_media_placeholder(
        doc,
        "video",
        "3–5 minute walkthrough: Teams install → SSO → host lobby → two joiners (Teams + web guest) → "
        "start game → one round → licensed host shows premium word set → optional meeting side panel.",
    )
    add_body(doc, "Suggested video URL (replace before submit): https://www.youtube.com/watch?v=REPLACE_ME")

    add_heading(doc, "Support contact", 1)
    add_body(
        doc,
        "Publisher: Annaliese Sebastian LLC. "
        "Support email for validation: REPLACE-WITH-SUPPORT-EMAIL@domain.com",
    )
    add_assumption(
        doc,
        "Replace support email with the address listed in Partner Center before submission.",
    )

    add_heading(doc, "Health check (for reviewers)", 1)
    add_body(
        doc,
        "API health: GET https://ingroups-production.up.railway.app/api/health — expect JSON with "
        "ok status and teamsSso: true when production is configured.",
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
